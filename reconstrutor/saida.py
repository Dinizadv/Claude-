"""Geração da planilha final multi-aba."""

from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .esquema import CONF_ALTA, CONF_BAIXA, CONF_MEDIA
from .pasta import CandidatoPasta


COLUNAS_SUGESTOES = [
    "ID origem",
    "Descrição original",
    "Descrição sugerida final",
    "Cliente sugerido",
    "Tipo de vínculo sugerido",
    "Número da pasta sugerida",
    "Nome da pasta sugerida",
    "Tipo/Subtipo sugerido",
    "Idioma sugerido",
    "Data de início",
    "Hora de início sugerida",
    "Tempo total resultante",
    "Duração original",
    "Executante",
    "Grau de confiança",
    "Necessita revisão humana?",
    "Indício de desdobramento?",
    "Observações",
    "Base comparativa principal utilizada",
    "Pasta alternativa plausível, se houver",
]


PREENCHIMENTOS = {
    CONF_ALTA:  PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid"),
    CONF_MEDIA: PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid"),
    CONF_BAIXA: PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid"),
}


def _texto_readme(caminho_a: Path, caminho_b: Path) -> list[str]:
    return [
        "RECONSTRUÇÃO DE LANÇAMENTOS DE HORAS — LEGAL ONE",
        "",
        f"Gerado em: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}",
        f"Arquivo do usuário: {caminho_a.name}",
        f"Arquivo da equipe:  {caminho_b.name}",
        "",
        "FINALIDADE",
        "Reconstruir, linha a linha, como cada lançamento deveria ser feito no Legal One,",
        "com base nos padrões reais de cliente, pasta, tipo/subtipo e idioma observados",
        "na base comparativa da equipe.",
        "",
        "REGRA TEMPORAL CRÍTICA",
        "Data de início e duração total do lançamento original são preservadas verbatim.",
        "A 'Hora de início sugerida' é 00:00, refletindo o fluxo real do Legal One",
        "(o sistema calcula data/hora final a partir da duração).",
        "",
        "LEGENDA — Grau de confiança",
        "  alta : cliente claro, pasta dominante, vínculo decisivo, descrição reescrita.",
        "  media: cliente identificado, pasta plausível com alternativa, vínculo razoável.",
        "  baixa: cliente presumido, pasta ausente, dúvida SERV/PROC ou descrição telegráfica.",
        "         Toda linha 'baixa' aparece também em CASOS_DUVIDOSOS.",
        "",
        "ABAS",
        "  README                : este texto.",
        "  RAW_USUARIO           : dados brutos do Arquivo A (preservados).",
        "  RAW_EQUIPE            : dados brutos do Arquivo B (preservados).",
        "  SUGESTOES_FINAIS      : 20 colunas operacionais, prontas para uso.",
        "  CASOS_DUVIDOSOS       : subset que exige revisão humana.",
        "  PASTAS_REFERENCIADAS  : pastas da equipe efetivamente consideradas.",
        "",
        "REGRAS NEGATIVAS APLICADAS",
        "  - nunca inventar número de pasta;",
        "  - nunca inventar processo ou cliente;",
        "  - nunca alterar data de início ou duração total sem inconsistência objetiva.",
    ]


def _ajustar_larguras(ws, larguras: dict[int, int] | None = None) -> None:
    larguras = larguras or {}
    for col_idx, col_cells in enumerate(ws.columns, start=1):
        if col_idx in larguras:
            ws.column_dimensions[get_column_letter(col_idx)].width = larguras[col_idx]
            continue
        max_len = 12
        for c in col_cells:
            v = c.value
            if v is None:
                continue
            max_len = max(max_len, min(60, len(str(v)) + 2))
        ws.column_dimensions[get_column_letter(col_idx)].width = max_len


def _escrever_dataframe(ws, df: pd.DataFrame) -> None:
    for col_idx, nome in enumerate(df.columns, start=1):
        cell = ws.cell(row=1, column=col_idx, value=str(nome))
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color="DDDDDD", end_color="DDDDDD", fill_type="solid")
    for row_idx, (_, row) in enumerate(df.iterrows(), start=2):
        for col_idx, val in enumerate(row, start=1):
            if isinstance(val, float) and pd.isna(val):
                val = None
            ws.cell(row=row_idx, column=col_idx, value=val)
    ws.freeze_panes = "A2"


def escrever(
    caminho_saida: Path,
    raw_usuario: pd.DataFrame,
    raw_equipe: pd.DataFrame,
    sugestoes: list[dict[str, Any]],
    indice_pastas: list[CandidatoPasta],
    caminho_a: Path,
    caminho_b: Path,
) -> None:
    wb = Workbook()

    ws_readme = wb.active
    ws_readme.title = "README"
    for i, linha in enumerate(_texto_readme(caminho_a, caminho_b), start=1):
        c = ws_readme.cell(row=i, column=1, value=linha)
        if i == 1:
            c.font = Font(bold=True, size=14)
        elif linha.isupper() and linha.strip():
            c.font = Font(bold=True)
        c.alignment = Alignment(wrap_text=False)
    ws_readme.column_dimensions["A"].width = 100

    ws_a = wb.create_sheet("RAW_USUARIO")
    _escrever_dataframe(ws_a, raw_usuario)
    _ajustar_larguras(ws_a)

    ws_b = wb.create_sheet("RAW_EQUIPE")
    _escrever_dataframe(ws_b, raw_equipe)
    _ajustar_larguras(ws_b)

    df_sug = pd.DataFrame(sugestoes, columns=COLUNAS_SUGESTOES)
    ws_sug = wb.create_sheet("SUGESTOES_FINAIS")
    _escrever_dataframe(ws_sug, df_sug)
    col_conf = COLUNAS_SUGESTOES.index("Grau de confiança") + 1
    for row_idx in range(2, len(df_sug) + 2):
        nivel = ws_sug.cell(row=row_idx, column=col_conf).value
        if nivel in PREENCHIMENTOS:
            ws_sug.cell(row=row_idx, column=col_conf).fill = PREENCHIMENTOS[nivel]
    _ajustar_larguras(ws_sug, larguras={2: 40, 3: 50, 18: 40})

    duvidosos = df_sug[
        (df_sug["Grau de confiança"] == CONF_BAIXA)
        | (df_sug["Necessita revisão humana?"] == True)
        | (df_sug["Indício de desdobramento?"] == True)
    ]
    ws_dub = wb.create_sheet("CASOS_DUVIDOSOS")
    _escrever_dataframe(ws_dub, duvidosos.reset_index(drop=True))
    _ajustar_larguras(ws_dub, larguras={2: 40, 3: 50, 18: 40})

    df_pastas = pd.DataFrame([
        {
            "Cliente": p.cliente,
            "Vínculo": p.vinculo,
            "Número da pasta": p.numero,
            "Nome da pasta": p.nome,
            "Frequência de uso": p.freq,
            "Idioma predominante": p.idioma,
            "Tipos/subtipos recorrentes": "; ".join(p.tipos_subtipos[:5]),
        }
        for p in sorted(indice_pastas, key=lambda c: (c.cliente, c.vinculo, -c.freq))
    ])
    ws_pas = wb.create_sheet("PASTAS_REFERENCIADAS")
    if df_pastas.empty:
        ws_pas.cell(row=1, column=1, value="(nenhuma pasta referenciada)")
    else:
        _escrever_dataframe(ws_pas, df_pastas)
        _ajustar_larguras(ws_pas)

    caminho_saida = Path(caminho_saida)
    caminho_saida.parent.mkdir(parents=True, exist_ok=True)
    wb.save(caminho_saida)
