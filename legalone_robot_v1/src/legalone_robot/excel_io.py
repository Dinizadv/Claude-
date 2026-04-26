
from __future__ import annotations

from datetime import datetime, time, timedelta
from pathlib import Path
from typing import List

from openpyxl import load_workbook, Workbook

from .models import LaunchRow


EXPECTED_COLUMNS = [
    "linha_id",
    "executante",
    "data_inicio",
    "hora_inicio",
    "duracao_hhmm",
    "cliente_principal",
    "negociacao",
    "descricao_negociacao",
    "pasta",
    "nome_pasta",
    "tipo_subtipo",
    "descricao",
    "cobravel",
    "observacoes_executante",
    "gerente_conta",
    "grupo",
    "pode_lancar",
    "motivo_bloqueio",
    "status_execucao",
    "id_lancamento_retorno",
    "mensagem_retorno",
]


def _as_string(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, time):
        return value.strftime("%H:%M")
    if isinstance(value, timedelta):
        total_seconds = int(value.total_seconds())
        hours = total_seconds // 3600
        minutes = (total_seconds % 3600) // 60
        return f"{hours:02d}:{minutes:02d}"
    return str(value).strip()


def load_rows(path: str) -> list[LaunchRow]:
    workbook = load_workbook(path)
    sheet = workbook[workbook.sheetnames[0]]
    headers = [_as_string(cell.value) for cell in sheet[1]]
    missing = [col for col in EXPECTED_COLUMNS if col not in headers]
    if missing:
        raise ValueError(f"Planilha sem colunas obrigatórias: {missing}")

    header_map = {name: idx + 1 for idx, name in enumerate(headers)}
    rows: list[LaunchRow] = []

    for row_idx in range(2, sheet.max_row + 1):
        raw = {col: _as_string(sheet.cell(row_idx, header_map[col]).value) for col in EXPECTED_COLUMNS}
        if not any(raw.values()):
            continue
        rows.append(LaunchRow(**raw))
    return rows


def write_results(original_path: str, output_path: str, rows: list[LaunchRow]) -> None:
    workbook = load_workbook(original_path)
    sheet = workbook[workbook.sheetnames[0]]
    headers = [_as_string(cell.value) for cell in sheet[1]]
    header_map = {name: idx + 1 for idx, name in enumerate(headers)}
    row_by_id = {str(sheet.cell(r, header_map["linha_id"]).value).strip(): r for r in range(2, sheet.max_row + 1)}

    for item in rows:
        excel_row = row_by_id.get(item.linha_id)
        if not excel_row:
            continue
        sheet.cell(excel_row, header_map["status_execucao"]).value = item.status_execucao
        sheet.cell(excel_row, header_map["id_lancamento_retorno"]).value = item.id_lancamento_retorno
        sheet.cell(excel_row, header_map["mensagem_retorno"]).value = item.mensagem_retorno
        sheet.cell(excel_row, header_map["motivo_bloqueio"]).value = item.motivo_bloqueio

    workbook.save(output_path)
