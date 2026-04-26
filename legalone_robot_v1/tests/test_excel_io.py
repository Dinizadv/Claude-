from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from legalone_robot.excel_io import EXPECTED_COLUMNS, load_rows, write_results
from legalone_robot.models import LaunchRow


def _make_workbook(tmp_path: Path, rows: list[dict], columns: list[str] | None = None) -> Path:
    cols = columns if columns is not None else EXPECTED_COLUMNS
    wb = Workbook()
    sheet = wb.active
    for col_idx, name in enumerate(cols, start=1):
        sheet.cell(1, col_idx).value = name
    for row_idx, row in enumerate(rows, start=2):
        for col_idx, name in enumerate(cols, start=1):
            sheet.cell(row_idx, col_idx).value = row.get(name, "")
    path = tmp_path / "input.xlsx"
    wb.save(path)
    return path


def test_load_rows_happy_path(tmp_path):
    path = _make_workbook(tmp_path, [
        {"linha_id": "1", "executante": "Joao", "data_inicio": "2026-04-26",
         "hora_inicio": "09:00", "duracao_hhmm": "01:30",
         "tipo_subtipo": "Consulta|Geral", "descricao": "ok"},
    ])
    rows = load_rows(str(path))
    assert len(rows) == 1
    assert rows[0].linha_id == "1"
    assert rows[0].executante == "Joao"
    assert rows[0].duracao_hhmm == "01:30"


def test_load_rows_skips_fully_blank_lines(tmp_path):
    path = _make_workbook(tmp_path, [
        {"linha_id": "1", "executante": "Joao", "data_inicio": "2026-04-26",
         "hora_inicio": "09:00", "duracao_hhmm": "01:30",
         "tipo_subtipo": "T|S", "descricao": "x"},
        {},
        {"linha_id": "2", "executante": "Maria", "data_inicio": "2026-04-26",
         "hora_inicio": "10:00", "duracao_hhmm": "00:30",
         "tipo_subtipo": "T|S", "descricao": "y"},
    ])
    rows = load_rows(str(path))
    assert [r.linha_id for r in rows] == ["1", "2"]


def test_load_rows_raises_when_required_columns_missing(tmp_path):
    columns = [c for c in EXPECTED_COLUMNS if c != "executante"]
    path = _make_workbook(tmp_path, [{"linha_id": "1"}], columns=columns)
    with pytest.raises(ValueError, match="executante"):
        load_rows(str(path))


def test_write_results_updates_status_columns(tmp_path):
    path = _make_workbook(tmp_path, [
        {"linha_id": "1", "executante": "Joao", "data_inicio": "2026-04-26",
         "hora_inicio": "09:00", "duracao_hhmm": "01:30",
         "tipo_subtipo": "T|S", "descricao": "x"},
        {"linha_id": "2", "executante": "Maria", "data_inicio": "2026-04-26",
         "hora_inicio": "10:00", "duracao_hhmm": "00:30",
         "tipo_subtipo": "T|S", "descricao": "y"},
    ])
    rows = load_rows(str(path))
    rows[0].status_execucao = "SUCESSO"
    rows[0].id_lancamento_retorno = "ID-101"
    rows[0].mensagem_retorno = "ok"
    rows[1].status_execucao = "BLOQUEADO"
    rows[1].motivo_bloqueio = "campo X ausente"
    rows[1].mensagem_retorno = "campo X ausente"

    out_path = tmp_path / "out.xlsx"
    write_results(str(path), str(out_path), rows)

    wb = load_workbook(out_path)
    sheet = wb.active
    headers = [cell.value for cell in sheet[1]]
    idx = {name: i + 1 for i, name in enumerate(headers)}

    assert sheet.cell(2, idx["status_execucao"]).value == "SUCESSO"
    assert sheet.cell(2, idx["id_lancamento_retorno"]).value == "ID-101"
    assert sheet.cell(2, idx["mensagem_retorno"]).value == "ok"
    assert sheet.cell(3, idx["status_execucao"]).value == "BLOQUEADO"
    assert sheet.cell(3, idx["motivo_bloqueio"]).value == "campo X ausente"
